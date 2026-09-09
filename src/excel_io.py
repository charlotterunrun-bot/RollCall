"""Excel parsing and transactional persistence for roll-call records."""

from __future__ import annotations

import hashlib
import math
import os
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
import xlrd

import paths
import storage
from errors import AppError

HEADERS = ["序号", "学号", "姓名", "班级"]
ENGLISH_HEADERS = ["No.", "Student ID", "Name", "Class"]
FIELD_ORDER = ("seq", "no", "name", "clazz")
FIELD_ALIASES = {
    "seq": {"序号", "no.", "no", "sequence"},
    "no": {"学号", "student id", "studentid"},
    "name": {"姓名", "name"},
    "clazz": {"班级", "class"},
}
STATUS_ALIASES = {"到": "到", "假": "假", "旷": "旷", "present": "到", "leave": "假", "absent": "旷"}
_DATE_YMD = re.compile(r"^(\d{4})([-/])(\d{1,2})\2(\d{1,2})$")
_DATE_DMY = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")
_DATE_LIKE = re.compile(r"\d{1,4}[-/]\d{1,2}[-/]\d{1,4}")
_ZERO_FORMAT = re.compile(r"^0+$")
_UNSET = object()


def _app_error(code: str, **params) -> AppError:
    return AppError(code, **params)


def _fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip().strip("\ufeff").strip()


def _to_int(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = _fmt(value)
    if re.fullmatch(r"[+-]?\d+", text):
        try:
            return int(text)
        except ValueError:
            return None
    return None


def _is_formula(cell) -> bool:
    # Literal strings beginning with '=' are data_type 's'; actual formulas are 'f'.
    return getattr(cell, "data_type", None) == "f"


def _header_key(value):
    text = _fmt(value)
    if not text:
        return None
    folded = text.casefold()
    for key, aliases in FIELD_ALIASES.items():
        if text in aliases or folded in {a.casefold() for a in aliases}:
            return key
    return None


def _normalize_status(value, *, path, row, column):
    text = _fmt(value)
    if not text:
        return ""
    result = STATUS_ALIASES.get(text.casefold())
    if result is None:
        raise _app_error("excel.invalid_status", value=text, path=str(path), row=row, column=column)
    return result


def _normalize_date_header(value, *, path, column):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _fmt(value)
    if not text:
        return None
    if _DATE_DMY.fullmatch(text):
        raise _app_error("excel.ambiguous_date_header", header=text, path=str(path), column=column)
    match = _DATE_YMD.fullmatch(text)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(3)), int(match.group(4))).isoformat()
        except ValueError as exc:
            raise _app_error("excel.invalid_date_header", header=text, path=str(path), column=column) from exc
    if _DATE_LIKE.search(text):
        raise _app_error("excel.invalid_date_header", header=text, path=str(path), column=column)
    return None


def _normalize_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _fmt(value)
    match = _DATE_YMD.fullmatch(text)
    if not match:
        raise _app_error("excel.invalid_date", date=text)
    try:
        return date(int(match.group(1)), int(match.group(3)), int(match.group(4))).isoformat()
    except ValueError as exc:
        raise _app_error("excel.invalid_date", date=text) from exc


def _number_text(cell, *, path, row, column, field):
    value = cell.value
    if value is None:
        return ""
    if field == "no" and isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and not value.is_integer():
            raise _app_error("excel.invalid_student_id", row=row, column=column, path=str(path))
        integer = int(value)
        digits = str(abs(integer))
        if len(digits) > 15:
            raise _app_error("excel.student_id_precision", row=row, column=column, path=str(path))
        number_format = str(getattr(cell, "number_format", ""))
        if _ZERO_FORMAT.fullmatch(number_format):
            width = len(number_format)
            return ("-" if integer < 0 else "") + digits.zfill(width)
        return str(integer)
    return _fmt(value)


class Student:
    __slots__ = ("row", "seq", "seq_num", "no", "name", "clazz", "records")

    def __init__(self, row, seq, no, name, clazz, records=None):
        self.row = row
        self.seq = seq
        self.seq_num = _to_int(seq)
        self.no = no
        self.name = name
        self.clazz = clazz
        self.records = records or {}

    @property
    def id(self):
        return self.no


class RollCallData:
    def __init__(self, students, date_columns, *, source_path=None, fingerprint=None, sheet_name=None, field_columns=None, storage_language="zh_CN", header_row=1):
        self.students = students
        self.source_path = Path(source_path) if source_path is not None else None
        self.fingerprint = fingerprint
        self.sheet_name = sheet_name
        self.field_columns = field_columns or {}
        self.date_columns = date_columns
        self.storage_language = storage_language
        self.header_row = header_row


def _read_stable_bytes(path: Path):
    try:
        first = path.read_bytes()
    except FileNotFoundError as exc:
        raise _app_error("storage_source_missing", path=str(path)) from exc
    except OSError as exc:
        raise _app_error("storage_read_failed", path=str(path)) from exc
    digest = hashlib.sha256(first).hexdigest()
    try:
        second = path.read_bytes()
    except OSError as exc:
        raise _app_error("storage_read_failed", path=str(path)) from exc
    if second != first:
        raise _app_error("storage_conflict", path=str(path))
    return first, digest


def _sheet_header(ws, path):
    max_column = max(1, ws.max_column)
    for row in range(1, ws.max_row + 1):
        values = [ws.cell(row=row, column=col).value for col in range(1, max_column + 1)]
        if not any(_fmt(value) for value in values):
            continue
        columns = {}
        for col in range(1, max_column + 1):
            cell = ws.cell(row=row, column=col)
            key = _header_key(cell.value)
            if key is not None:
                if _is_formula(cell):
                    raise _app_error("excel.formula_key_field", path=str(path), row=row, column=col)
                if key in columns:
                    raise _app_error("excel.duplicate_field", field=key, path=str(path), column=col)
                columns[key] = col
        if set(columns) == set(FIELD_ORDER):
            return row, columns
    return None


def _select_sheet(wb, path, sheet_name):
    if sheet_name is not None:
        if sheet_name not in wb.sheetnames:
            raise _app_error("excel.sheet_not_found", sheet=sheet_name, path=str(path))
        selected = _sheet_header(wb[sheet_name], path)
        if selected is None:
            raise _app_error("excel.invalid_header", sheet=sheet_name, path=str(path))
        return wb[sheet_name], selected
    record_name = next((name for name in wb.sheetnames if name.casefold() == "record"), None)
    if record_name is not None:
        record_header = _sheet_header(wb[record_name], path)
        if record_header is not None:
            return wb[record_name], record_header
    matches = []
    for name in wb.sheetnames:
        if name == record_name:
            continue
        try:
            selected = _sheet_header(wb[name], path)
        except AppError:
            # An unrelated sheet may contain an incomplete or duplicate alias
            # set; only a selected record sheet can veto the workbook.
            continue
        if selected is not None:
            matches.append((name, selected))
    if not matches:
        raise _app_error("excel.invalid_header", path=str(path))
    if len(matches) != 1:
        raise _app_error("excel.ambiguous_sheets", sheets=[name for name, _ in matches], path=str(path))
    return wb[matches[0][0]], matches[0][1]


def _merged_key_area(ws, columns, date_columns, header_row, path):
    key_columns = set(columns.values()) | set(date_columns.values())
    for merged in ws.merged_cells.ranges:
        if any(col in key_columns for col in range(merged.min_col, merged.max_col + 1)) and merged.max_row >= header_row:
            raise _app_error("excel.merged_key_field", path=str(path), range=str(merged))


def _parse_workbook(wb, path, sheet_name=None):
    ws, (header_row, field_columns) = _select_sheet(wb, path, sheet_name)
    for col in range(1, max(1, ws.max_column) + 1):
        if _is_formula(ws.cell(row=header_row, column=col)):
            raise _app_error("excel.formula_key_field", path=str(path), row=header_row, column=col)
    max_column = max(1, ws.max_column)
    date_columns = {}
    for col in range(1, max_column + 1):
        cell = ws.cell(row=header_row, column=col)
        value = _normalize_date_header(cell.value, path=path, column=col)
        if value is not None:
            if value in date_columns:
                raise _app_error("excel.duplicate_date", date=value, columns=[date_columns[value], col], path=str(path))
            date_columns[value] = col
    _merged_key_area(ws, field_columns, date_columns, header_row, path)
    chinese_headers = {"序号", "学号", "姓名", "班级"}
    english = all(
        _fmt(ws.cell(header_row, field_columns[key]).value).casefold()
        in {alias.casefold() for alias in FIELD_ALIASES[key] if alias not in chinese_headers}
        for key in FIELD_ORDER
    )
    storage_language = "en_US" if english else "zh_CN"
    has_chinese_status = False
    students = []
    seen = set()
    for row in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row=row, column=col).value for col in range(1, max_column + 1)]
        if not any(_fmt(value) for value in values):
            continue
        for col in field_columns.values():
            if _is_formula(ws.cell(row=row, column=col)):
                raise _app_error("excel.formula_key_field", path=str(path), row=row, column=col)
        seq_col = field_columns["seq"]
        seq_cell = ws.cell(row=row, column=seq_col)
        seq = _number_text(seq_cell, path=path, row=row, column=seq_col, field="seq")
        if seq_cell.value is not None and isinstance(seq_cell.value, (float, int)) and not isinstance(seq_cell.value, bool) and _to_int(seq_cell.value) is None:
            raise _app_error("excel.invalid_sequence", path=str(path), row=row, column=seq_col)
        if isinstance(seq_cell.value, str):
            text = seq_cell.value.strip()
            if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
                try:
                    numeric = float(text)
                except ValueError:
                    numeric = math.inf
                if not math.isfinite(numeric) or not numeric.is_integer():
                    raise _app_error("excel.invalid_sequence", path=str(path), row=row, column=seq_col)
                seq = str(int(numeric))
        no_col = field_columns["no"]
        no = _number_text(ws.cell(row=row, column=no_col), path=path, row=row, column=no_col, field="no")
        name = _fmt(ws.cell(row=row, column=field_columns["name"]).value)
        clazz = _fmt(ws.cell(row=row, column=field_columns["clazz"]).value)
        if not no:
            raise _app_error("excel.empty_student_id", path=str(path), row=row)
        if not name:
            raise _app_error("excel.empty_student_name", path=str(path), row=row)
        if no in seen:
            raise _app_error("excel.duplicate_student_id", no=no, path=str(path), row=row)
        seen.add(no)
        records = {}
        for normalized_date, col in date_columns.items():
            cell = ws.cell(row=row, column=col)
            if _is_formula(cell):
                raise _app_error("excel.formula_key_field", path=str(path), row=row, column=col)
            status = _normalize_status(cell.value, path=path, row=row, column=col)
            if _fmt(cell.value) in {"到", "假", "旷"}:
                has_chinese_status = True
            if status:
                records[normalized_date] = status
        students.append(Student(row, seq, no, name, clazz, records))
    if has_chinese_status:
        storage_language = "zh_CN"
    return RollCallData(students, date_columns, source_path=path, sheet_name=ws.title, field_columns=field_columns, storage_language=storage_language, header_row=header_row)


def load_record(path=None, *, sheet_name=None) -> RollCallData:
    target = Path(path if path is not None else paths.record_path())
    raw, digest = _read_stable_bytes(target)
    try:
        wb = load_workbook(BytesIO(raw), data_only=False)
    except Exception as exc:
        raise _app_error("excel.invalid_file", path=str(target)) from exc
    try:
        data = _parse_workbook(wb, target, sheet_name)
    finally:
        wb.close()
    data.fingerprint = digest
    if storage.fingerprint(target) != digest:
        raise _app_error("storage_conflict", path=str(target))
    return data


def is_first_run() -> bool:
    if not os.path.isfile(paths.record_path()):
        return True
    try:
        data = load_record()
    except Exception:
        return True
    return len(data.students) == 0


def _namelist_header(values, path, row):
    columns = {}
    for column, value in enumerate(values, start=1):
        key = _header_key(value)
        if key is None:
            continue
        if key in columns:
            raise _app_error("excel.duplicate_field", field=key, path=str(path), row=row, column=column)
        columns[key] = column
    return columns if set(columns) == set(FIELD_ORDER) else None


def _parse_namelist_cells(rows, columns, path, *, header_row, get_cell, merged_ranges=()):
    key_columns = set(columns.values())
    for min_row, max_row, min_col, max_col in merged_ranges:
        if max_row >= header_row and any(col in key_columns for col in range(min_col, max_col + 1)):
            raise _app_error("excel.merged_key_field", path=str(path), range=f"{min_row}:{max_row},{min_col}:{max_col}")
    students = []
    seen = set()
    for row in range(header_row + 1, len(rows) + 1):
        cells = [get_cell(row, col) for col in range(1, max(columns.values()) + 1)]
        if not any(_fmt(value) for value in cells):
            continue
        for key, col in columns.items():
            if _is_formula(get_cell(row, col, with_meta=True)):
                raise _app_error("excel.formula_key_field", path=str(path), row=row, column=col)
        seq_col = columns["seq"]
        seq_cell = get_cell(row, seq_col, with_meta=True)
        seq = _sequence_text(seq_cell, path=path, row=row, column=seq_col)
        no_col = columns["no"]
        no = _number_text(get_cell(row, no_col, with_meta=True), path=path, row=row, column=no_col, field="no")
        name = _fmt(get_cell(row, columns["name"]))
        clazz = _fmt(get_cell(row, columns["clazz"]))
        if not no:
            raise _app_error("excel.empty_student_id", path=str(path), row=row)
        if not name:
            raise _app_error("excel.empty_student_name", path=str(path), row=row)
        if no in seen:
            raise _app_error("excel.duplicate_student_id", no=no, path=str(path), row=row)
        seen.add(no)
        students.append({"seq": seq, "no": no, "name": name, "clazz": clazz})
    if not students:
        raise _app_error("excel.empty_namelist", path=str(path))
    return students


def _sequence_text(cell, *, path, row, column):
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(float(value)) or not float(value).is_integer():
            raise _app_error("excel.invalid_sequence", path=str(path), row=row, column=column)
        return str(int(value))
    text = _fmt(value)
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
        try:
            numeric = float(text)
        except ValueError:
            numeric = math.inf
        if not math.isfinite(numeric) or not numeric.is_integer():
            raise _app_error("excel.invalid_sequence", path=str(path), row=row, column=column)
        return str(int(numeric))
    return text


def _read_namelist_xlsx(path, *, sheet_name=None):
    target = Path(path)
    try:
        wb = load_workbook(target, data_only=False, read_only=False)
    except Exception as exc:
        raise _app_error("excel.invalid_file", path=str(target)) from exc
    try:
        ws, (header_row, columns) = _select_sheet(wb, target, sheet_name)
        for col in range(1, ws.max_column + 1):
            if _is_formula(ws.cell(header_row, col)):
                raise _app_error("excel.formula_key_field", path=str(target), row=header_row, column=col)
        values = [[ws.cell(row=row, column=col).value for col in range(1, max(1, ws.max_column) + 1)] for row in range(1, ws.max_row + 1)]
        def get_cell(row, col, with_meta=False):
            return ws.cell(row=row, column=col) if with_meta else ws.cell(row=row, column=col).value
        merged = [(r.min_row, r.max_row, r.min_col, r.max_col) for r in ws.merged_cells.ranges]
        return _parse_namelist_cells(values, columns, target, header_row=header_row, get_cell=get_cell, merged_ranges=merged)
    finally:
        wb.close()


def _xls_formula_cells(raw):
    try:
        from xlrd.compdoc import CompDoc
        stream, offset, length = CompDoc(raw, logfile=None).locate_named_stream("Workbook")
        stream = stream[offset:offset + length]
    except Exception as exc:
        raise _app_error("excel.invalid_file") from exc
    formulas = []
    pos = 0
    sheet_index = -1
    while pos + 4 <= len(stream):
        opcode = int.from_bytes(stream[pos:pos + 2], "little")
        size = int.from_bytes(stream[pos + 2:pos + 4], "little")
        payload = stream[pos + 4:pos + 4 + size]
        if len(payload) != size:
            break
        if opcode in {0x0809, 0x0009, 0x0209, 0x0409} and len(payload) >= 4:
            stream_type = int.from_bytes(payload[2:4], "little")
            if stream_type in {0x0010, 0x0020, 0x0040}:
                sheet_index += 1
        if opcode in {0x0006, 0x0206, 0x0406} and len(payload) >= 6:
            formulas.append((sheet_index, int.from_bytes(payload[0:2], "little") + 1, int.from_bytes(payload[2:4], "little") + 1))
        pos += 4 + size
    return formulas


def _xls_cell(book, sheet, row, col):
    value = sheet.cell_value(row - 1, col - 1)
    number_format = ""
    try:
        xf = book.xf_list[sheet.cell_xf_index(row - 1, col - 1)]
        number_format = book.format_map[xf.format_key].format_str
    except (IndexError, KeyError, AttributeError):
        pass
    return type("XlsCell", (), {"value": value, "data_type": "s", "number_format": number_format})()


def _read_namelist_xls(path, *, sheet_name=None):
    target = Path(path)
    try:
        raw = target.read_bytes()
        formulas = _xls_formula_cells(raw)
        book = xlrd.open_workbook(file_contents=raw, formatting_info=True)
    except AppError:
        raise
    except Exception as exc:
        raise _app_error("excel.invalid_file", path=str(target)) from exc
    try:
        matches = []
        for index, sheet in enumerate(book.sheets()):
            for row in range(1, sheet.nrows + 1):
                values = [sheet.cell_value(row - 1, col - 1) for col in range(1, sheet.ncols + 1)]
                try:
                    columns = _namelist_header(values, target, row)
                except AppError:
                    if sheet.name.casefold() == "record" or sheet.name == sheet_name:
                        raise
                    continue
                if columns is not None:
                    matches.append((index, row, columns))
                    break
        if sheet_name is not None:
            selected = next((entry for entry in matches if book.sheet_by_index(entry[0]).name == sheet_name), None)
            if selected is None:
                raise _app_error("excel.sheet_not_found", sheet=sheet_name, path=str(target))
        else:
            record = next((entry for entry in matches if book.sheet_by_index(entry[0]).name.casefold() == "record"), None)
            if record is not None:
                selected = record
            elif len(matches) == 1:
                selected = matches[0]
            elif len(matches) > 1:
                raise _app_error("excel.ambiguous_sheets", sheets=[book.sheet_by_index(e[0]).name for e in matches], path=str(target))
            else:
                raise _app_error("excel.invalid_header", path=str(target))
        index, header_row, columns = selected
        critical_formula = any(
            sheet_id == -1 or (sheet_id == index and row >= header_row and column in columns.values())
            for sheet_id, row, column in formulas
        )
        if critical_formula:
            raise _app_error("excel.formula_key_field", path=str(target))
        sheet = book.sheet_by_index(index)
        rows = [[sheet.cell_value(row, col) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
        def get_cell(row, col, with_meta=False):
            return _xls_cell(book, sheet, row, col) if with_meta else sheet.cell_value(row - 1, col - 1)
        merged = [(rlo + 1, rhi, clo + 1, chi) for rlo, rhi, clo, chi in sheet.merged_cells]
        return _parse_namelist_cells(rows, columns, target, header_row=header_row, get_cell=get_cell, merged_ranges=merged)
    finally:
        book.release_resources()


def read_namelist(path, *, sheet_name=None) -> list:
    suffix = Path(path).suffix.casefold()
    if suffix == ".xls":
        return _read_namelist_xls(path, sheet_name=sheet_name)
    if suffix == ".xlsx":
        return _read_namelist_xlsx(path, sheet_name=sheet_name)
    raise _app_error("excel.unsupported_namelist", path=str(path))


def _set_literal(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def _status_for_language(value, language):
    if language.casefold() in {"en", "en_us", "english"}:
        return {"到": "Present", "假": "Leave", "旷": "Absent"}[value]
    return value


def _record_bytes(students, language):
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    headers = ENGLISH_HEADERS if language.casefold() in {"en", "en_us", "english"} else HEADERS
    for col, value in enumerate(headers, start=1):
        _set_literal(ws.cell(row=1, column=col), value)
    for row, student in enumerate(students, start=2):
        for col, value in enumerate([student["seq"], student["no"], student["name"], student["clazz"]], start=1):
            _set_literal(ws.cell(row=row, column=col), value)
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def create_record_from_namelist(students, *, language="zh_CN", path=None, expected_fingerprint=_UNSET):
    target = Path(path if path is not None else paths.record_path())
    if expected_fingerprint is _UNSET:
        expected_fingerprint = storage.fingerprint(target)
    raw = _record_bytes(students, language)
    candidate = {}

    def writer(temp):
        temp.write_bytes(raw)

    def validator(temp):
        candidate["data"] = load_record(temp)

    committed = storage.atomic_write(target, writer, validator, expected_fingerprint=expected_fingerprint, backup_kind="auto")
    data = candidate["data"]
    data.source_path = target
    data.fingerprint = committed
    return data


def _copy_and_update(raw, data: RollCallData, no, normalized_date, normalized_value, path):
    try:
        wb = load_workbook(BytesIO(raw), data_only=False)
    except Exception as exc:
        raise _app_error("excel.invalid_file", path=str(path)) from exc
    try:
        if data.sheet_name not in wb.sheetnames:
            raise _app_error("excel.sheet_not_found", sheet=data.sheet_name, path=str(path))
        ws = wb[data.sheet_name]
        row = next((student.row for student in data.students if student.id == no), None)
        if row is None:
            raise _app_error("excel.student_not_found", no=no, path=str(path))
        col = data.date_columns.get(normalized_date)
        if col is None:
            col = ws.max_column + 1
            _set_literal(ws.cell(row=data.header_row, column=col), normalized_date)
        _set_literal(ws.cell(row=row, column=col), normalized_value)
        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()
    finally:
        wb.close()


def write_record(no: str, date: str, value: str, *, data: RollCallData | None = None) -> RollCallData:
    if data is None:
        data = load_record()
    target = Path(data.source_path if data.source_path is not None else paths.record_path())
    normalized_date = _normalize_date(date)
    normalized_value = _normalize_status(value, path=target, row=0, column=0)
    if not normalized_value:
        raise _app_error("excel.invalid_status", value=value, path=str(target))
    student = next((item for item in data.students if item.id == _fmt(no)), None)
    if student is None:
        raise _app_error("excel.student_not_found", no=_fmt(no), path=str(target))
    if normalized_date in student.records and student.records[normalized_date]:
        raise _app_error("excel.duplicate_attendance", no=student.id, date=normalized_date, path=str(target))
    raw, digest = _read_stable_bytes(target)
    if data.fingerprint != digest:
        raise _app_error("storage_conflict", path=str(target))
    candidate = {}

    def writer(temp):
        temp.write_bytes(_copy_and_update(raw, data, student.id, normalized_date, _status_for_language(normalized_value, data.storage_language), target))

    def validator(temp):
        candidate["data"] = load_record(temp, sheet_name=data.sheet_name)
        refreshed = next((item for item in candidate["data"].students if item.id == student.id), None)
        if refreshed is None or refreshed.records.get(normalized_date) != normalized_value:
            raise _app_error("excel.write_validation_failed", no=student.id, date=normalized_date)

    committed = storage.atomic_write(target, writer, validator, expected_fingerprint=data.fingerprint, backup_kind="auto")
    result = candidate["data"]
    result.source_path = target
    result.fingerprint = committed
    return result
