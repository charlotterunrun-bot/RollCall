"""Bounded, offscreen smoke flow used only by packaged native builds."""

from __future__ import annotations

import json
import struct
import traceback
from pathlib import Path


def run_smoke(data_dir: str | Path) -> int:
    """Exercise the bundled Qt UI and workbook transaction in a fresh folder."""
    import paths

    target_dir = Path(data_dir).expanduser().resolve()
    if not target_dir.is_dir() or any(target_dir.iterdir()):
        return 2
    paths.set_data_dir(target_dir)
    report = target_dir / "smoke-result.json"
    def progress(stage):
        report.write_text(json.dumps({"mode": "packaged-offscreen", "stage": stage}, indent=2), encoding="utf-8")
    try:
        progress("imports")
        from PySide6.QtCore import QSysInfo, QTimer
        from PySide6.QtWidgets import QApplication

        import excel_io
        import i18n
        from init_window import InitWindow
        from main_window import MainWindow
        from openpyxl import load_workbook
        from version import __version__

        progress("qt")
        app = QApplication.instance() or QApplication(["RollCall", "--smoke-test"])
        app.setQuitOnLastWindowClosed(False)
        i18n.set_language("zh_CN")
        for template in ("template-zh.xlsx", "template-en.xlsx"):
            template_path = Path(paths.resource_path(f"resources/{template}"))
            workbook = load_workbook(template_path, read_only=True)
            try:
                assert workbook.active.max_row == 1
            finally:
                workbook.close()
        progress("template")
        students = [{"seq": "1", "no": "SMOKE-1", "name": "Smoke One", "clazz": "A"}, {"seq": "2", "no": "SMOKE-2", "name": "Smoke Two", "clazz": "A"}]
        record = target_dir / "record.xlsx"
        from openpyxl import Workbook
        roster = target_dir / "smoke-roster.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "学号", "姓名", "班级"])
        for student in students:
            sheet.append([student["seq"], student["no"], student["name"], student["clazz"]])
        workbook.save(roster)
        workbook.close()
        init = InitWindow(target_path=record)
        init.show()
        app.processEvents()
        if __version__ not in init.windowTitle():
            raise AssertionError("version missing from initialization title")
        imported = init.import_roster_path(roster)
        if imported is None or len(imported) != len(students):
            raise AssertionError("first-run roster import did not commit")
        init.close()
        roster.unlink(missing_ok=True)
        data = excel_io.load_record(record)
        progress("initial-record")
        first = MainWindow(data=data, data_path=record, acquire_lock=True)
        first.marquee_enabled = True
        first.marquee_duration = 1
        first.show()
        first.start_rollcall()
        first._marquee_tick()
        QTimer.singleShot(30, app.quit)
        app.exec()
        progress("first-reveal")
        if first._reveal_timer.isActive() or not first.btn_present.isEnabled():
            raise AssertionError("Qt reveal timer did not deliver")
        first.on_record("到")
        first.close()

        i18n.set_language("en_US")
        resumed = MainWindow(data=excel_io.load_record(record), data_path=record, acquire_lock=True)
        resumed.retranslate_ui()
        resumed.marquee_duration = 1
        resumed.show()
        resumed.start_rollcall()
        QTimer.singleShot(30, app.quit)
        app.exec()
        progress("resume-reveal")
        if resumed._reveal_timer.isActive() or not resumed.btn_present.isEnabled():
            raise AssertionError("Qt resume reveal timer did not deliver")
        resumed.on_record("Present")
        resumed.close()

        completed = MainWindow(data=excel_io.load_record(record), data_path=record, acquire_lock=True)
        completed.start_rollcall()
        if not completed._finished or completed.current is not None:
            raise AssertionError("a completed day started a second round")
        completed.close()
        # Let Qt deliver one real event-loop turn while the windows are closed.
        QTimer.singleShot(0, app.quit)
        app.exec()
        reloaded = excel_io.load_record(record)
        actual_records = {student.id: dict(student.records) for student in reloaded.students}
        payload = {
            "mode": "packaged-offscreen",
            "platform": __import__("sys").platform,
            "arch": QSysInfo.currentCpuArchitecture(),
            "pointer_bits": struct.calcsize("P") * 8,
            "frozen": bool(getattr(__import__("sys"), "frozen", False)),
            "version": __version__,
            "record_reloaded": excel_io.probe_record(record) == "ready",
            "languages": ["zh_CN", "en_US"],
            "attendance_count": sum(len(records) for records in actual_records.values()),
            "students": sorted(actual_records),
            "statuses": sorted(status for records in actual_records.values() for status in records.values()),
        }
        report.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        return 0
    except Exception as exc:
        report.write_text(json.dumps({"mode": "packaged-offscreen", "error": str(exc), "traceback": traceback.format_exc()}, indent=2), encoding="utf-8")
        return 1
