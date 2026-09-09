"""Focused tests for Excel structure discovery and persistence."""

from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from errors import AppError
from excel_io import create_record_from_namelist, load_record, write_record


def _save(path, sheets):
    wb = Workbook()
    first = wb.active
    for index, (title, rows) in enumerate(sheets):
        ws = first if index == 0 else wb.create_sheet(title)
        ws.title = title
        for row in rows:
            ws.append(row)
    wb.save(path)
    wb.close()


def test_empty_column_does_not_shift_dates(tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["序号", "学号", "姓名", "班级", "2026-09-08", None, "2026-09-09"], [1, "S001", "Example", "A", "到", None, "假"]])])

    data = load_record(path=p)

    assert data.students[0].records == {"2026-09-08": "到", "2026-09-09": "假"}
    assert data.date_columns["2026-09-09"] == 7
    assert data.field_columns == {"seq": 1, "no": 2, "name": 3, "clazz": 4}
    assert data.sheet_name == "record"


@pytest.mark.parametrize(
    "header, expected_language, status",
    [
        (["No.", "Student ID", "Name", "Class", "2026-09-09"], "en_US", "Present"),
        ([" sequence ", "student id", "name", "class", "2026-09-09"], "en_US", "Present"),
        (["序号", "Student ID", "姓名", "Class", "2026-09-09"], "zh_CN", "到"),
    ],
)
def test_aliases_and_storage_language(tmp_path, header, expected_language, status):
    p = tmp_path / "record.xlsx"
    _save(p, [("other", [["ignore"], ["x"]]), ("sheet", [header, ["1", "0007", "Ada", "A", status]])])
    data = load_record(p)
    assert data.sheet_name == "sheet"
    assert data.storage_language == expected_language
    assert data.students[0].no == "0007"
    assert data.date_columns == {"2026-09-09": 5}


def test_multiple_matching_sheets_require_selection(tmp_path):
    p = tmp_path / "record.xlsx"
    header = ["序号", "学号", "姓名", "班级", "2026-09-09"]
    _save(p, [("one", [header, [1, "S1", "One", "A", None]]), ("two", [header, [1, "S2", "Two", "B", None]])])

    with pytest.raises(AppError) as caught:
        load_record(p)
    assert caught.value.code == "excel.ambiguous_sheets"
    assert caught.value.params["sheets"] == ["one", "two"]
    assert load_record(p, sheet_name="two").students[0].no == "S2"


@pytest.mark.parametrize("header", [["序号", "学号", "姓名", "班级", "2026-02-30"], ["序号", "学号", "姓名", "班级", "09/10/2026"]])
def test_invalid_or_ambiguous_date_headers_are_rejected(tmp_path, header):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [header, [1, "S1", "One", "A", None]])])
    with pytest.raises(AppError) as caught:
        load_record(p)
    assert caught.value.code in {"excel.invalid_date_header", "excel.ambiguous_date_header"}


def test_formula_date_header_is_rejected_but_unrelated_sheet_is_ignored(tmp_path):
    p = tmp_path / "record.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append(["序号", "学号", "姓名", "班级", "2026-09-09"])
    ws.append([1, "S1", "One", "A", None])
    notes = wb.create_sheet("notes")
    notes.append(["序号", "学号", "姓名", "班级"])
    notes.append([1, "S2", "Two", "B"])
    ws["E1"] = "=DATE(2026,9,9)"
    wb.save(p)
    wb.close()
    with pytest.raises(AppError) as caught:
        load_record(p)
    assert caught.value.code in {"excel.formula_key_field", "excel.invalid_header"}

    ws = Workbook().active
    ws.title = "record"
    ws.append(["序号", "学号", "姓名", "班级"])
    ws.append([1, "S1", "One", "A"])
    # A malformed duplicate-field sheet must not veto a valid record sheet.
    wb2 = Workbook()
    valid = wb2.active
    valid.title = "record"
    valid.append(["序号", "学号", "姓名", "班级"])
    valid.append([1, "S1", "One", "A"])
    bad = wb2.create_sheet("notes")
    bad.append(["学号", "学号", "姓名", "班级"])
    bad.append(["S2", "S2", "Two", "B"])
    wb2.save(tmp_path / "valid_with_notes.xlsx")
    wb2.close()
    assert load_record(tmp_path / "valid_with_notes.xlsx").sheet_name == "record"


def test_datetime_and_reordered_fields_are_normalized(tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["备注", "姓名", "2026/9/8", "班级", "学号", "序号"], ["keep", "Ada", "到", "A", "S1", 2]])])
    data = load_record(p)
    assert data.field_columns == {"seq": 6, "no": 5, "name": 2, "clazz": 4}
    assert data.date_columns == {"2026-09-08": 3}
    assert data.students[0].name == "Ada"
    assert data.students[0].records == {"2026-09-08": "到"}

    p2 = tmp_path / "datetime.xlsx"
    _save(p2, [("record", [["序号", "学号", "姓名", "班级", datetime(2026, 9, 9)], [1, "S2", "B", "A", "假"]])])
    assert load_record(p2).date_columns == {"2026-09-09": 5}


def test_leading_blank_rows_keep_header_row_for_writes(tmp_path):
    p = tmp_path / "leading_blank.xlsx"
    _save(p, [("record", [[None] * 5, [None] * 5, ["序号", "学号", "姓名", "班级", "备注"], [1, "S1", "One", "A", "keep"]])])
    data = load_record(p)
    assert data.header_row == 3
    write_record("S1", "2026-09-09", "到", data=data)
    wb = load_workbook(p, data_only=False)
    assert wb["record"]["F3"].value == "2026-09-09"
    assert wb["record"]["F4"].value == "到"
    assert wb["record"]["A1"].value is None
    wb.close()


@pytest.mark.parametrize("rows, code", [
    ([["序号", "学号", "姓名", "班级", "2026-09-09"], [1, "S1", "One", "A", "Present"]], None),
    ([["序号", "学号", "姓名", "班级", "2026-09-09"], [1, "S1", "One", "A", "Leave"]], None),
    ([["序号", "学号", "姓名", "班级", "2026-09-09"], [1, "S1", "One", "A", "Absent"]], None),
    ([["序号", "学号", "姓名", "班级", "2026-09-09"], [1, "S1", "One", "A", "Mystery"]], "excel.invalid_status"),
])
def test_status_aliases_are_normalized(rows, code, tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", rows)])
    if code:
        with pytest.raises(AppError) as caught:
            load_record(p)
        assert caught.value.code == code
    else:
        assert set(load_record(p).students[0].records.values()) <= {"到", "假", "旷"}


def test_english_headers_with_chinese_statuses_use_chinese_storage(tmp_path):
    p = tmp_path / "mixed.xlsx"
    _save(p, [("record", [["No.", "Student ID", "Name", "Class", "2026-09-09"], [1, "S1", "One", "A", "到"]])])
    data = load_record(p)
    assert data.storage_language == "zh_CN"


def test_english_record_writes_english_status(tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["No.", "Student ID", "Name", "Class", "2026-09-09"], [1, "S1", "One", "A", None]])])
    data = load_record(p)
    write_record("S1", "2026-09-09", "到", data=data)
    wb = load_workbook(p, data_only=False)
    assert wb["record"]["E2"].value == "Present"
    wb.close()


def test_noninteger_numeric_sequence_text_is_rejected(tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["序号", "学号", "姓名", "班级"], ["1.5", "S1", "One", "A"]])])
    with pytest.raises(AppError) as caught:
        load_record(p)
    assert caught.value.code == "excel.invalid_sequence"


def test_formula_in_key_field_and_duplicate_values_are_rejected(tmp_path):
    p = tmp_path / "record.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append(["序号", "学号", "姓名", "班级", "2026-09-09"])
    ws.append([1, "S1", "One", "A", None])
    ws.append([2, "S1", "Two", "B", None])
    ws["B2"] = "=A2"
    wb.save(p)
    wb.close()
    with pytest.raises(AppError) as caught:
        load_record(p)
    assert caught.value.code in {"excel.formula_key_field", "excel.duplicate_student_id"}


def test_merged_key_area_is_rejected_but_literal_equals_value_is_allowed(tmp_path):
    p = tmp_path / "record.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append(["序号", "学号", "姓名", "班级", "2026-09-09"])
    ws.append([1, "=S1", "One", "A", None])
    ws.merge_cells("A2:A3")
    wb.save(p)
    wb.close()
    with pytest.raises(AppError) as caught:
        load_record(p)
    assert caught.value.code == "excel.merged_key_field"

    q = tmp_path / "literal.xlsx"
    literal = Workbook()
    literal.active.title = "record"
    literal.active.append(["序号", "学号", "姓名", "班级"])
    literal.active.append([1, "=S1", "One", "A"])
    literal.active["B2"].data_type = "s"
    literal.save(q)
    literal.close()
    assert load_record(q).students[0].no == "=S1"


def test_numeric_student_id_zero_mask_and_precision_loss(tmp_path):
    p = tmp_path / "masked.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append(["序号", "学号", "姓名", "班级"])
    ws.append([1, 7, "One", "A"])
    ws["B2"].number_format = "00000"
    wb.save(p)
    wb.close()
    assert load_record(p).students[0].no == "00007"

    q = tmp_path / "precision.xlsx"
    _save(q, [("record", [["序号", "学号", "姓名", "班级"], [1, 1234567890123456, "One", "A"]])])
    with pytest.raises(AppError) as caught:
        load_record(q)
    assert caught.value.code == "excel.student_id_precision"


def test_corrupt_and_empty_files_are_errors(tmp_path):
    corrupt = tmp_path / "bad.xlsx"
    corrupt.write_bytes(b"not an xlsx")
    with pytest.raises(AppError) as caught:
        load_record(corrupt)
    assert caught.value.code == "excel.invalid_file"

    empty = tmp_path / "empty.xlsx"
    _save(empty, [("record", [["序号", "学号", "姓名", "班级"]])])
    assert load_record(empty).students == []


def test_duplicate_date_and_noninteger_sequence_are_rejected(tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["序号", "学号", "姓名", "班级", "2026-09-09", "2026/9/9"], [1.5, "S1", "One", "A", None, None]])])
    with pytest.raises(AppError) as caught:
        load_record(p)
    assert caught.value.code in {"excel.duplicate_date", "excel.invalid_sequence"}


def test_write_is_transactional_preserves_sheets_and_reloads_snapshot(tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["序号", "学号", "姓名", "班级", "备注", "2026-09-09"], [1, "S1", "One", "A", "keep", None]]), ("notes", [["untouched"]])])
    data = load_record(p)
    written = write_record("S1", "2026-09-09", "到", data=data)
    assert written.students[0].records == {"2026-09-09": "到"}
    assert written.fingerprint
    wb = load_workbook(p, data_only=False)
    assert wb.sheetnames == ["record", "notes"]
    assert wb["record"]["E2"].value == "keep"
    assert wb["notes"]["A1"].value == "untouched"
    wb.close()


def test_write_returns_committed_candidate_when_final_reload_would_fail(tmp_path, monkeypatch):
    from pathlib import Path

    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["序号", "学号", "姓名", "班级"], [1, "S1", "One", "A"]])])
    data = load_record(p)
    excel_module = __import__("excel_io")
    original_load = excel_module.load_record

    def fail_only_final_target(path=None, *, sheet_name=None):
        if path is not None and Path(path) == p:
            raise RuntimeError("simulated postcommit reload failure")
        return original_load(path, sheet_name=sheet_name)

    monkeypatch.setattr(excel_module, "load_record", fail_only_final_target)
    result = write_record("S1", "2026-09-09", "到", data=data)

    assert result.source_path == p
    assert result.fingerprint
    assert result.students[0].records == {"2026-09-09": "到"}
    assert p.is_file()


def test_write_rejects_second_attendance_and_preserves_bytes(tmp_path):
    p = tmp_path / "record.xlsx"
    _save(p, [("record", [["序号", "学号", "姓名", "班级", "2026-09-09"], [1, "S1", "One", "A", "到"]])])
    data = load_record(p)
    before = p.read_bytes()
    with pytest.raises(AppError) as caught:
        write_record("S1", "2026-09-09", "假", data=data)
    assert caught.value.code == "excel.duplicate_attendance"
    assert p.read_bytes() == before


def test_create_record_uses_language_and_keeps_literal_equals_text(tmp_path, monkeypatch):
    import paths

    monkeypatch.setattr(paths, "record_path", lambda: str(tmp_path / "record.xlsx"))
    monkeypatch.setattr(paths, "record_dir", lambda: str(tmp_path))
    create_record_from_namelist([{"seq": "1", "no": "=S1", "name": "=Ada", "clazz": "A"}], language="en_US")
    wb = load_workbook(tmp_path / "record.xlsx", data_only=False)
    assert [cell.value for cell in wb.active[1][:4]] == ["No.", "Student ID", "Name", "Class"]
    assert wb.active["B2"].data_type == "s"
    assert wb.active["B2"].value == "=S1"
    wb.close()
