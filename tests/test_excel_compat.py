"""Compatibility regressions for legacy record files."""

from openpyxl import Workbook, load_workbook

from excel_io import load_record, read_namelist, write_record


def test_legacy_chinese_record_writes_chinese_status(tmp_path):
    path = tmp_path / "legacy.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append([" 序号\ufeff", "学号", "姓名", "班级", "备注", "2026-09-08"])
    ws.append([1, "0001", "旧记录", "A", "保留", "到"])
    wb.save(path)
    wb.close()

    data = load_record(path)
    result = write_record("0001", "2026-09-09", "假", data=data)

    assert result.storage_language == "zh_CN"
    assert result.students[0].records == {"2026-09-08": "到", "2026-09-09": "假"}
    reopened = load_workbook(path, data_only=False)
    assert [cell.value for cell in reopened["record"][1][:7]] == [" 序号\ufeff", "学号", "姓名", "班级", "备注", "2026-09-08", "2026-09-09"]
    assert [cell.value for cell in reopened["record"][2][:7]] == [1, "0001", "旧记录", "A", "保留", "到", "假"]
    reopened.close()


def test_snapshot_conflict_stops_write(tmp_path):
    path = tmp_path / "record.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append(["序号", "学号", "姓名", "班级"])
    ws.append([1, "S1", "One", "A"])
    wb.save(path)
    wb.close()
    data = load_record(path)

    changed = Workbook()
    changed.remove(changed.active)
    sheet = changed.create_sheet("record")
    sheet.append(["序号", "学号", "姓名", "班级"])
    sheet.append([1, "S1", "Two", "A"])
    changed.save(path)
    changed.close()

    import pytest
    from errors import AppError

    with pytest.raises(AppError) as caught:
        write_record("S1", "2026-09-09", "到", data=data)
    assert caught.value.code == "storage_conflict"


def test_xlsx_namelist_uses_aliases_physical_columns_and_zero_mask(tmp_path):
    path = tmp_path / "namelist.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "roster"
    ws.append(["备注", "Name", "Student ID", "Sequence", "Class"])
    ws.append(["keep", "One", 7, "2.0", "A"])
    ws["C2"].number_format = "00000"
    wb.save(path)
    wb.close()

    assert read_namelist(path) == [{"seq": "2", "no": "00007", "name": "One", "clazz": "A"}]


def test_xlsx_namelist_sheet_ambiguity_and_formula_are_structured(tmp_path):
    path = tmp_path / "ambiguous.xlsx"
    wb = Workbook()
    first = wb.active
    first.title = "one"
    first.append(["序号", "学号", "姓名", "班级"])
    first.append([1, "S1", "One", "A"])
    second = wb.create_sheet("two")
    second.append(["序号", "学号", "姓名", "班级"])
    second.append([2, "S2", "Two", "B"])
    wb.save(path)
    wb.close()
    import pytest
    from errors import AppError
    with pytest.raises(AppError) as caught:
        read_namelist(path)
    assert caught.value.code == "excel.ambiguous_sheets"
    assert read_namelist(path, sheet_name="two")[0]["no"] == "S2"

    formula_path = tmp_path / "formula.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "record"
    ws.append(["序号", "学号", "姓名", "班级"])
    ws.append(["=1+1", "S1", "One", "A"])
    wb.save(formula_path)
    wb.close()
    with pytest.raises(AppError) as caught:
        read_namelist(formula_path)
    assert caught.value.code == "excel.formula_key_field"


def test_xls_namelist_preserves_zero_mask_and_rejects_formula(tmp_path):
    xlwt = __import__("xlwt")
    path = tmp_path / "legacy.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("record")
    for col, value in enumerate(["序号", "学号", "姓名", "班级"]):
        sheet.write(0, col, value)
    sheet.write(1, 0, 2.0)
    sheet.write(1, 1, 7, xlwt.easyxf(num_format_str="00000"))
    sheet.write(1, 2, "One")
    sheet.write(1, 3, "A")
    workbook.save(str(path))
    assert read_namelist(path) == [{"seq": "2", "no": "00007", "name": "One", "clazz": "A"}]

    formula_path = tmp_path / "formula.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("record")
    for col, value in enumerate(["序号", "学号", "姓名", "班级"]):
        sheet.write(0, col, value)
    sheet.write(1, 0, xlwt.Formula("1+1"))
    sheet.write(1, 1, "S1")
    sheet.write(1, 2, "One")
    sheet.write(1, 3, "A")
    workbook.save(str(formula_path))
    import pytest
    from errors import AppError
    with pytest.raises(AppError) as caught:
        read_namelist(formula_path)
    assert caught.value.code == "excel.formula_key_field"

    remark_path = tmp_path / "remark_formula.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("record")
    for col, value in enumerate(["序号", "学号", "姓名", "班级", "备注"]):
        sheet.write(0, col, value)
    sheet.write(1, 0, 1)
    sheet.write(1, 1, "S1")
    sheet.write(1, 2, "One")
    sheet.write(1, 3, "A")
    sheet.write(1, 4, xlwt.Formula("1+1"))
    workbook.save(str(remark_path))
    assert read_namelist(remark_path)[0]["no"] == "S1"
