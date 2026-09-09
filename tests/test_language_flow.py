from datetime import date
from pathlib import Path

import pytest


@pytest.fixture(autouse=True)
def reset_language():
    import i18n
    i18n.set_language("zh_CN")
    yield
    i18n.set_language("zh_CN")


@pytest.fixture
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _data():
    from excel_io import RollCallData, Student
    return RollCallData([Student(2, "1", "S1", "A", {}, {})], {}, source_path="record.xlsx", sheet_name="record")


def _two_student_data():
    from excel_io import RollCallData, Student
    return RollCallData([
        Student(2, "1", "S1", "A", {}, {}),
        Student(3, "2", "S2", "A", {}, {}),
    ], {}, source_path="record.xlsx", sheet_name="record")


def test_main_window_retranslates_without_changing_student_or_timers(qapp):
    import i18n
    from main_window import MainWindow

    i18n.set_language("zh_CN")
    win = MainWindow(data=_data(), acquire_lock=False)
    win.current = win.data.students[0]
    win._marquee_timer.start()
    before = (win.current, win._marquee_timer.isActive(), win._reveal_timer.isActive(), win._marquee_timer.timerId(), win._reveal_timer.timerId())
    i18n.set_language("en_US")
    win.retranslate_ui()
    assert win.current is before[0]
    assert win._marquee_timer.isActive() == before[1]
    assert win._reveal_timer.isActive() == before[2]
    assert win._marquee_timer.timerId() == before[3]
    assert win._reveal_timer.timerId() == before[4]
    assert win.btn_present.text() == "Present"
    win.close()


def test_language_switch_during_real_marquee_and_reveal_preserves_state(qapp):
    import i18n
    from main_window import MainWindow

    i18n.set_language("zh_CN")
    win = MainWindow(data=_two_student_data(), acquire_lock=False)
    win.start_rollcall()
    assert win._marquee_timer.isActive()
    target = win.current
    win._marquee_tick()
    marquee_text = (win.no_label.text(), win.name_label.text())
    i18n.set_language("en_US")
    win.retranslate_ui()
    assert win.current is target
    assert win._marquee_timer.isActive()
    assert win._reveal_timer.isActive()
    assert (win.no_label.text(), win.name_label.text()) == marquee_text
    win._reveal()
    revealed = (win.no_label.text(), win.name_label.text())
    i18n.set_language("zh_CN")
    win.retranslate_ui()
    assert (win.no_label.text(), win.name_label.text()) == revealed
    assert win.current is target
    win.close()


def test_language_switch_on_completed_page_only_retranslates_completion(qapp, monkeypatch):
    import i18n
    import main_window
    from main_window import MainWindow
    from excel_io import RollCallData, Student

    monkeypatch.setattr(main_window, "current_date_string", lambda: "2026-09-09")
    data = RollCallData([Student(2, "1", "S1", "A", {}, {"2026-09-09": "到"})], {}, source_path="record.xlsx", sheet_name="record")
    win = MainWindow(data=data, acquire_lock=False)
    win.start_rollcall()
    assert win._finished
    i18n.set_language("en_US")
    win.retranslate_ui()
    assert win.name_label.text() == "Today's roll call is complete!"
    assert win.current is None
    win.close()


def test_language_save_failure_keeps_active_language_student_and_timers(qapp, monkeypatch):
    import i18n
    from errors import AppError
    from main_window import MainWindow

    win = MainWindow(data=_two_student_data(), acquire_lock=False)
    win.start_rollcall()
    target = win.current
    timer_state = (win._marquee_timer.isActive(), win._reveal_timer.isActive(), win._marquee_timer.timerId(), win._reveal_timer.timerId())
    monkeypatch.setattr("main_window.config.save_settings", lambda **kwargs: (_ for _ in ()).throw(AppError("config.write_failed")))
    monkeypatch.setattr(i18n, "critical", lambda *args, **kwargs: None)
    i18n.set_language("zh_CN")
    win._change_language("en_US")
    assert i18n.language() == "zh_CN"
    assert win.current is target
    assert (win._marquee_timer.isActive(), win._reveal_timer.isActive(), win._marquee_timer.timerId(), win._reveal_timer.timerId()) == timer_state
    win.close()


def test_init_window_uses_active_language_for_ui_and_new_roster(qapp, tmp_path, monkeypatch):
    import i18n
    from init_window import InitWindow

    i18n.set_language("en_US")
    window = InitWindow(target_path=tmp_path / "record.xlsx")
    assert window.windowTitle().startswith("RollCall")
    assert window.btn_pick.text() == "Choose roster file"
    captured = {}
    monkeypatch.setattr(window, "_read_students", lambda path: [{"seq": "1", "no": "S1", "name": "Ada", "clazz": "A"}])
    monkeypatch.setattr("init_window.QFileDialog.getOpenFileName", lambda *args, **kwargs: (str(tmp_path / "roster.xlsx"), ""))
    monkeypatch.setattr("init_window.excel_io.create_record_from_namelist", lambda students, **kwargs: captured.update(kwargs))
    monkeypatch.setattr(i18n, "information", lambda *args, **kwargs: None)
    window.on_pick_namelist()
    assert captured["language"] == "en_US"
    window.close()


def test_english_geometry_keeps_window_and_config_readable(qapp):
    import i18n
    from config_dialog import ConfigDialog
    from main_window import MainWindow

    i18n.set_language("en_US")
    main = MainWindow(data=_two_student_data(), acquire_lock=False)
    main.resize(640, 480)
    assert main.minimumSize().width() >= 640
    assert main.minimumSize().height() >= 480
    assert main.btn_present.sizeHint().width() < main.width() / 2
    dialog = ConfigDialog(current="random_count")
    assert dialog.minimumWidth() >= 520
    assert dialog.sizeHint().width() >= dialog.minimumWidth()
    dialog.close()
    main.close()


def test_config_dialog_has_language_and_translated_strategy_labels(qapp):
    import i18n
    from config_dialog import ConfigDialog

    i18n.set_language("en_US")
    dialog = ConfigDialog(current="random_count")
    assert dialog.windowTitle() == "Roll-call settings"
    assert dialog.language_selected() == "en_US"
    assert "historical" in dialog._radios["random_count"].text().lower()
    dialog.close()


def test_recovery_dialog_retranslates_error_and_standard_buttons(qapp):
    import i18n
    from errors import AppError
    from recovery_dialog import RecoveryDialog

    i18n.set_language("en_US")
    dialog = RecoveryDialog("record.xlsx", AppError("excel.invalid_file", path="record.xlsx"))
    dialog.retranslate_ui()
    assert dialog.windowTitle() == "Record recovery"
    assert dialog.retry_btn.text() == "Retry"
    assert "Excel" in dialog.detail.text()
    dialog.close()


def test_empty_templates_use_matching_standard_headers(tmp_path):
    from openpyxl import load_workbook
    import paths

    for name, expected in (
        ("template-zh.xlsx", ("序号", "学号", "姓名", "班级")),
        ("template-en.xlsx", ("No.", "Student ID", "Name", "Class")),
    ):
        path = Path(paths.resource_path(f"resources/{name}"))
        workbook = load_workbook(path, read_only=True, data_only=False)
        assert tuple(next(workbook.active.iter_rows(max_row=1, values_only=True))) == expected
        assert workbook.active.max_row == 1
        workbook.close()
