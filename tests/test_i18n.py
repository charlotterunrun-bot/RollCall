import json
from pathlib import Path

import pytest


@pytest.fixture(autouse=True)
def reset_language():
    import i18n
    i18n.set_language("zh_CN")
    yield
    i18n.set_language("zh_CN")


def test_locale_resources_have_matching_keys_and_placeholders():
    import i18n

    zh = json.loads(Path(i18n.locale_path("zh_CN")).read_text(encoding="utf-8"))
    en = json.loads(Path(i18n.locale_path("en_US")).read_text(encoding="utf-8"))
    assert set(zh) == set(en)
    assert {k for k in zh if "{" in zh[k]} == {k for k in en if "{" in en[k]}
    for key in zh:
        assert i18n.placeholders(zh[key]) == i18n.placeholders(en[key])


def test_translation_and_unknown_language_fallback():
    import i18n

    i18n.set_language("zh_CN")
    assert i18n.tr("app.title") == "课堂点名"
    i18n.set_language("unsupported")
    assert i18n.language() == "en_US"
    assert i18n.tr("app.title") == "RollCall"
    assert i18n.tr("missing.key") == "missing.key"


def test_error_translation_tolerates_optional_context():
    import i18n
    from errors import AppError

    i18n.set_language("en_US")
    assert "attendance status" in i18n.error_text(AppError("excel.invalid_status"))
    assert "settings" in i18n.error_text(AppError("config.invalid", field="strategy"))


def test_language_persistence_validates_types(tmp_path, monkeypatch):
    import config
    import paths

    monkeypatch.setattr(paths, "record_dir", lambda: str(tmp_path))
    config.save_settings(language="en_US")
    assert config.load_language() == "en_US"
    Path(paths.config_path()).write_text(json.dumps({"language": 1}), encoding="utf-8")
    with pytest.raises(Exception) as exc:
        config.load_language()
    assert getattr(exc.value, "code", None) == "config.invalid"


def test_initialize_uses_system_language_when_no_persisted_choice(monkeypatch, tmp_path):
    import i18n
    import paths

    monkeypatch.setattr(paths, "config_path", lambda: str(tmp_path / "missing-config.json"))
    class FakeLocale:
        def name(self):
            return "en_GB"

        def uiLanguages(self):
            return ["en-GB"]

    monkeypatch.setattr(i18n.QLocale, "system", staticmethod(lambda: FakeLocale()))
    assert i18n.initialize() is None
    assert i18n.language() == "en_US"


def test_initialize_maps_qt_chinese_locale_and_unknown_persisted_choice(monkeypatch, tmp_path):
    import i18n
    import paths

    class ChineseLocale:
        def name(self):
            return "zh_CN"

        def uiLanguages(self):
            return ["zh-Hans-CN", "zh-CN"]

    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps({"language": "pirate"}), encoding="utf-8")
    monkeypatch.setattr(paths, "config_path", lambda: str(config_path))
    monkeypatch.setattr(i18n.QLocale, "system", staticmethod(lambda: ChineseLocale()))
    assert i18n.initialize() is None
    assert i18n.language() == "en_US"


def test_initialize_preserves_system_language_for_legacy_config(monkeypatch, tmp_path):
    import i18n
    import paths

    class EnglishLocale:
        def name(self):
            return "en_GB"

        def uiLanguages(self):
            return ["en-GB"]

    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps({"strategy": "random_count", "marquee": True, "marquee_duration": 500}), encoding="utf-8")
    monkeypatch.setattr(paths, "config_path", lambda: str(config_path))
    monkeypatch.setattr(i18n.QLocale, "system", staticmethod(lambda: EnglishLocale()))
    assert i18n.initialize() is None
    assert i18n.language() == "en_US"


def test_error_context_keeps_supplied_actionable_details():
    import i18n
    from errors import AppError

    i18n.set_language("en_US")
    message = i18n.error_text(AppError("excel.invalid_status", value="late", path="C:/records/class.xlsx", row=12, column=8))
    assert all(value in message for value in ("late", "C:/records/class.xlsx", "12", "8"))
    config_message = i18n.error_text(AppError("config.invalid", field="language", path="C:/records/config.json", reason="invalid JSON"))
    assert all(value in config_message for value in ("language", "C:/records/config.json", "invalid JSON"))


def test_duplicate_date_error_localizes_physical_columns_from_parser(tmp_path):
    from openpyxl import Workbook
    import excel_io
    import i18n
    from errors import AppError

    path = tmp_path / "duplicate-date.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "record"
    sheet.append(["序号", "学号", "姓名", "班级", "2026-09-09", "备注", "2026/09/09"])
    sheet.append([1, "S1", "One", "A", "", "keep", ""])
    workbook.save(path)
    workbook.close()

    caught = None
    try:
        excel_io.load_record(path)
    except AppError as error:
        assert error.code == "excel.duplicate_date"
        caught = error
    else:
        raise AssertionError("duplicate date was accepted")

    for language, expected in (("en_US", "physical columns: 5, 7"), ("zh_CN", "物理列：5、7")):
        i18n.set_language(language)
        message = i18n.error_text(caught)
        assert "2026-09-09" in message
        assert expected in message
        assert "5" in message and "7" in message


def test_formula_error_location_is_optional_and_rendered_once():
    import i18n
    from errors import AppError

    i18n.set_language("en_US")
    without_location = i18n.error_text(AppError("excel.formula_key_field", path="C:/records/class.xlsx"))
    assert "row" not in without_location and "column" not in without_location
    assert "C:/records/class.xlsx" in without_location
    with_location = i18n.error_text(AppError("excel.formula_key_field", path="C:/records/class.xlsx", row=12, column=8))
    assert with_location.count("row: 12") == 1
    assert with_location.count("column: 8") == 1


def test_known_config_reasons_are_localized_and_unknown_reasons_preserved():
    import i18n
    from errors import AppError

    for language, expected in (("zh_CN", "无效 JSON"), ("en_US", "invalid JSON")):
        i18n.set_language(language)
        message = i18n.error_text(AppError("config.invalid", path="config.json", reason="invalid JSON"))
        assert expected in message
        if language == "zh_CN":
            assert "invalid JSON" not in message
    i18n.set_language("zh_CN")
    assert "custom reason" in i18n.error_text(AppError("config.invalid", reason="custom reason"))


def test_locale_json_has_no_duplicate_keys():
    import i18n

    for language in i18n.SUPPORTED_LANGUAGES:
        duplicate_keys = []

        def collect(pairs):
            seen = set()
            result = {}
            for key, value in pairs:
                if key in seen:
                    duplicate_keys.append(key)
                seen.add(key)
                result[key] = value
            return result

        json.loads(Path(i18n.locale_path(language)).read_text(encoding="utf-8"), object_pairs_hook=collect)
        assert duplicate_keys == []


def test_message_box_standard_buttons_use_active_language(qapp, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    import i18n

    seen = []

    def fake_exec(box):
        seen.append({box.standardButton(button): button.text() for button in box.buttons()})
        return int(QMessageBox.StandardButton.No)

    monkeypatch.setattr(QMessageBox, "exec", fake_exec)
    i18n.set_language("en_US")
    assert i18n.question(None, "Confirm", "Continue?", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No) == QMessageBox.StandardButton.No
    assert set(seen[-1].values()) == {"Yes", "No"}
    i18n.set_language("zh_CN")
    i18n.information(None, "完成", "完成")
    assert seen[-1].get(QMessageBox.StandardButton.Ok) == "确定"
